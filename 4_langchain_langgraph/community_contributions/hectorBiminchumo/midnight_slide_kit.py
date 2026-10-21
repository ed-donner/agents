"""Minimal midnight blue slide template: dark blue background, big type, one sky blue accent."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

MIDNIGHT = RGBColor(0x0B, 0x1D, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xA9, 0xC1, 0xE8)
SKY = RGBColor(0x5A, 0xA9, 0xFF)

FONT = "Helvetica Neue"


def _textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    return box


def build_midnight_slide(title: str, key_points: list[str], recommendation: str, outfile: str) -> None:
    """Build the one-slide midnight blue recommendation deck and save it to outfile."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = MIDNIGHT

    # Big centered headline with lots of air around it
    _textbox(slide, 1.0, 0.9, 11.33, 1.4, title, 54, WHITE, bold=True, align=PP_ALIGN.CENTER)

    # Short sky blue rule under the headline
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.07), Inches(2.45), Inches(1.2), Inches(0.05))
    rule.fill.solid()
    rule.fill.fore_color.rgb = SKY
    rule.line.fill.background()

    # Three short points, centered, in pale blue
    for i, point in enumerate(key_points[:3]):
        _textbox(slide, 1.0, 3.0 + i * 0.85, 11.33, 0.7, point, 26, PALE, align=PP_ALIGN.CENTER)

    # One sky blue sentence at the bottom
    _textbox(slide, 1.5, 5.9, 10.33, 1.0, recommendation, 22, SKY, bold=True, align=PP_ALIGN.CENTER)

    prs.save(outfile)
